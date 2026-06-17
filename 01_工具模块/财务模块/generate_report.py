import json
import os
import subprocess
import re
import urllib.request
from datetime import datetime
from docx import Document

# ===================== 配置区域 =====================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
_PROJ_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DATA_FILE = os.path.join(BASE_DIR, "local_archive.json")
OUTPUT_FOLDER = os.path.join(_PROJ_ROOT, "每日财务")

# 从核心配置读取模型和 Ollama 连接信息
_CONFIG_PATH = os.path.join(_PROJ_ROOT, "00_核心主体", "模型.json")
try:
    with open(_CONFIG_PATH, "r", encoding="utf-8") as f:
        _GLOBAL_CONFIG = json.load(f)
except Exception:
    _GLOBAL_CONFIG = {}
MODEL_NAME = _GLOBAL_CONFIG.get("model_name_small", "qwen2.5:3b")
OLLAMA_HOST = _GLOBAL_CONFIG.get("ollama_host", None)
OLLAMA_BIN = _GLOBAL_CONFIG.get("ollama_bin", None) or "ollama"


def _ollama_generate(prompt):
    """统一的 Ollama 调用：优先 HTTP API（WSL 兼容），回退 CLI 子进程"""
    if OLLAMA_HOST:
        url = f"{OLLAMA_HOST.rstrip('/')}/api/generate"
        body = json.dumps({"model": MODEL_NAME, "prompt": prompt, "stream": False},
                          ensure_ascii=False).encode("utf-8")
        req = urllib.request.Request(url, data=body,
                                     headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                data = json.loads(resp.read().decode("utf-8"))
                return data.get("response", "")
        except Exception:
            pass

    try:
        result = subprocess.run(
            [OLLAMA_BIN, "run", MODEL_NAME],
            input=prompt, text=True, capture_output=True, encoding="utf-8", timeout=60
        )
        return result.stdout.strip()
    except Exception as e:
        return f"总结生成失败：{e}"
# ===================================================

def clean_text(text):
    """过滤掉 Word 无法识别的控制字符"""
    if not isinstance(text, str):
        return ""
    return re.sub(r'[\x00-\x1f\x7f-\x9f]', '', text)

def load_data():
    try:
        with open(DATA_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except:
        return []

def generate_report():
    data = load_data()
    if not data:
        print("📭 数据为空，请先记账或导入数据。")
        return

    if not os.path.exists(OUTPUT_FOLDER):
        os.makedirs(OUTPUT_FOLDER)

    print("🤖 正在调用本地模型生成总结...")
    summary_prompt = f"请根据以下财务数据，生成一份简洁的财务总结（包含总支出、总收入、结余和主要趋势）。数据：{json.dumps(data, ensure_ascii=False)}"
    raw_text = _ollama_generate(summary_prompt)
    summary_text = clean_text(raw_text)

    doc = Document()
    doc.add_heading('📊 每日财务报告', 0)
    doc.add_paragraph(f'报告生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M")}')
    doc.add_paragraph()

    doc.add_heading('一、财务总结', level=1)
    for para in summary_text.split('\n'):
        if para.strip():
            # 写入时再次确保清洁
            doc.add_paragraph(clean_text(para.strip()))

    doc.add_heading('二、明细数据', level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Grid Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = '日期'
    hdr_cells[1].text = '类型'
    hdr_cells[2].text = '金额 (元)'
    hdr_cells[3].text = '来源'

    for item in data:
        row_cells = table.add_row().cells
        row_cells[0].text = clean_text(item.get('date', ''))
        row_cells[1].text = '收入' if item.get('type') == 'income' else '支出'
        row_cells[2].text = clean_text(str(item.get('amount', 0)))
        row_cells[3].text = clean_text(item.get('source', ''))

    filename = f"财务报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    full_path = os.path.join(OUTPUT_FOLDER, filename)
    doc.save(full_path)

    print("=" * 50)
    print(f"✅ 报告生成完毕！")
    print(f"📁 保存位置：{full_path}")
    print("=" * 50)

if __name__ == "__main__":
    generate_report()