# -*- coding: utf-8 -*-
"""v3.13.19 财务报告生成自检：Phase 1（基础聚合）+ Phase 2（分类统计/期间对比/明细）+ Phase 2.5（默认范围=今日/近3天 + scope 参数）

运行：.venv/Scripts/python.exe -X utf8 _selftest_v3.13.19.py
（需在项目根目录下，依赖 .venv 内 python-docx）
"""
import os
import re
import sys
import tempfile
from datetime import datetime, timedelta

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "01_工具模块", "财务模块"))

import finance_module

passed = 0
failed = 0


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"  PASS  {name}")
    else:
        failed += 1
        print(f"  FAIL  {name}  {detail}")


def _has_emoji(s):
    """检测字符串是否含常见 emoji 字符（Word/WPS 显示为方块的场景）。"""
    emoji_pat = re.compile(
        "[\U0001F000-\U0001FAFF"
        "\U00002600-\U000027BF"
        "\U00002190-\U000021FF"
        "\U00002B00-\U00002BFF]"
    )
    return bool(emoji_pat.search(s))


print("=" * 60)
print("[自检] v3.13.19 generate_report Phase 1 + Phase 2 + Phase 2.5（聚合 + 分类/对比/明细 + 默认范围与 scope）")
print("=" * 60)

# ---------- 1. _compute_report_summary 聚合计算 ----------
print("\n--- 1. 报告汇总计算（收入/支出/结余）---")
SAMPLE = [
    {"id": "a", "type": "expense", "date": "2026-08-25", "amount": 10, "source": "饭"},
    {"id": "b", "type": "expense", "date": "2026-08-25", "amount": 20, "source": "奶茶"},
    {"id": "c", "type": "income", "date": "2026-08-25", "amount": 100, "source": "工资"},
    {"id": "d", "type": "income", "date": "2026-08-25", "amount": 50.5, "source": "红包"},
]
summary = finance_module._compute_report_summary(SAMPLE)
check("收入合计 = 150.5", summary.get("income_total") == 150.5, summary)
check("支出合计 = 30", summary.get("expense_total") == 30, summary)
check("结余 = 收入 - 支出 = 120.5", summary.get("balance") == 120.5, summary)

# 空数据 → 全 0，不抛异常
empty_summary = finance_module._compute_report_summary([])
check("空数据不抛异常且全为 0",
      empty_summary == {"income_total": 0, "expense_total": 0, "balance": 0},
      empty_summary)

# ---------- 2. generate_report 生成 Word 文档 ----------
print("\n--- 2. generate_report Word 文档生成 ---")
with tempfile.TemporaryDirectory() as tmp:
    result = finance_module.generate_report(data=SAMPLE, output_dir=tmp)
    check("返回 status == success", result.get("status") == "success", result)
    path = result.get("file_path", "")
    check("返回 file_path 且文件已生成", bool(path) and os.path.exists(path), path)
    check("文件名符合 财务报告_YYYYMMDD_HHMMSS.docx",
          re.match(r"财务报告_\d{8}_\d{6}\.docx$", os.path.basename(path)) is not None,
          os.path.basename(path))

    from docx import Document
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    joined = "\n".join(texts)
    check("首段标题为「每日财务报告」（无 emoji）",
          texts and "每日财务报告" in texts[0] and not _has_emoji(texts[0]),
          texts[0] if texts else "(无段落)")
    check("报告含生成时间", bool(re.search(r"报告生成时间[:：]\s*\d{4}-\d{2}-\d{2}", joined)), joined)
    check("报告含收入合计 150.5", "150.5" in joined, joined)
    check("报告含支出合计 30", "30" in joined, joined)
    check("报告含结余 120.5", "120.5" in joined, joined)
    check("整篇文档无 emoji（Word 方块兼容）", not _has_emoji(joined), joined)

# ---------- 3. 空数据边界 ----------
print("\n--- 3. 空数据边界 ---")
with tempfile.TemporaryDirectory() as tmp:
    empty_result = finance_module.generate_report(data=[], output_dir=tmp)
    check("空数据返回 error 状态", empty_result.get("status") == "error", empty_result)
    check("空数据不生成文件",
          not empty_result.get("file_path") or not os.path.exists(empty_result.get("file_path", "")),
          empty_result)

# ---------- 4. _compute_category_breakdown 分类统计 ----------
print("\n--- 4. 分类统计（category 分组 + 占比）---")
CAT_SAMPLE = [
    {"id": "e1", "type": "expense", "date": "2026-08-25", "amount": 150.5, "source": "吃饭", "category": "餐饮"},
    {"id": "e2", "type": "expense", "date": "2026-08-25", "amount": 80, "source": "打车", "category": "交通"},
    {"id": "e3", "type": "expense", "date": "2026-08-25", "amount": 20, "source": "奶茶", "category": "饮品"},
    {"id": "i1", "type": "income", "date": "2026-08-25", "amount": 500, "source": "工资", "category": "收入"},
]
cat_break = finance_module._compute_category_breakdown(CAT_SAMPLE)
check("仅统计支出记录（收入不计入）", "收入" not in cat_break, list(cat_break))
check("按金额降序排列", list(cat_break.keys()) == ["餐饮", "交通", "饮品"], list(cat_break.keys()))
check("餐饮 150.5 占比 60%", cat_break.get("餐饮", {}).get("amount") == 150.5
      and cat_break.get("餐饮", {}).get("percentage") == 60, cat_break)
check("交通 80 占比 32%", cat_break.get("交通", {}).get("amount") == 80
      and cat_break.get("交通", {}).get("percentage") == 32, cat_break)
check("饮品 20 占比 8%", cat_break.get("饮品", {}).get("amount") == 20
      and cat_break.get("饮品", {}).get("percentage") == 8, cat_break)
check("空数据返回空 dict", finance_module._compute_category_breakdown([]) == {}, "")

# 占比合计 100
_pct_sum = sum(v["percentage"] for v in cat_break.values())
check("占比合计为 100%", _pct_sum == 100, f"sum={_pct_sum}")

# ---------- 5. _get_period_comparison 期间对比 ----------
print("\n--- 5. 期间对比（本周vs上周 / 本月vs上月）---")
FIXED_NOW = datetime(2026, 8, 25)  # 周二：本周 08-24~08-30，上月 07-01~07-31
COMP_SAMPLE = [
    {"id": "w1", "type": "expense", "date": "2026-08-25", "amount": 100},   # 本周+本月
    {"id": "w2", "type": "expense", "date": "2026-08-20", "amount": 40},    # 上周
    {"id": "m1", "type": "expense", "date": "2026-08-05", "amount": 200},   # 本月（非本周）
    {"id": "m2", "type": "expense", "date": "2026-07-15", "amount": 100},   # 上月
    {"id": "inc", "type": "income", "date": "2026-08-25", "amount": 9999},  # 收入不计入开销
]
comp = finance_module._get_period_comparison(COMP_SAMPLE, now=FIXED_NOW)
check("本周支出 = 100", comp["week"]["this_total"] == 100, comp["week"])
check("上周支出 = 40", comp["week"]["last_total"] == 40, comp["week"])
check("本周vs上周变化 +150%", comp["week"]["percent"] == 150, comp["week"])
# 本月 = 8月全部支出：本周100 + 本月200 + 上周40（08-20 属 8 月） = 340
check("本月支出 = 340（本周100 + 本月200 + 上周40均在8月）",
      comp["month"]["this_total"] == 340, comp["month"])
check("上月支出 = 100", comp["month"]["last_total"] == 100, comp["month"])
check("本月vs上月变化 +240%", comp["month"]["percent"] == 240, comp["month"])
check("收入不计入开销（inc 9999 未计入）",
      comp["week"]["this_total"] == 100 and comp["month"]["this_total"] == 340, comp)

# 上期为 0 → percent None（显示 "—"，不抛异常）
edge_comp = finance_module._get_period_comparison(
    [{"id": "x", "type": "expense", "date": "2026-08-25", "amount": 50}], now=FIXED_NOW)
check("上期为0时 percent 为 None", edge_comp["week"]["percent"] is None, edge_comp["week"])
check("None 格式化为 '—'", finance_module._format_percent_change(None) == "—", "")

# ---------- 6. _build_detail_table 交易明细排序与条数限制 ----------
print("\n--- 6. 交易明细（日期降序 + 最多50条）---")
DETAIL_SAMPLE = [
    {"id": "old", "type": "expense", "date": "2026-07-01", "time": "12:00"},
    {"id": "new", "type": "expense", "date": "2026-08-24", "time": "09:00"},
    {"id": "mid", "type": "expense", "date": "2026-08-01", "time": "12:00"},
]
detail = finance_module._build_detail_table(DETAIL_SAMPLE)
check("按日期降序（最新在前）", [r["id"] for r in detail] == ["new", "mid", "old"], [r["id"] for r in detail])
check("空数据返回空列表", finance_module._build_detail_table([]) == [], "")
detail_60 = finance_module._build_detail_table(DETAIL_SAMPLE * 20, limit=50)
check("60 条输入被限制到 50 条", len(detail_60) == 50, len(detail_60))
detail_2 = finance_module._build_detail_table(DETAIL_SAMPLE, limit=2)
check("limit=2 只返回 2 条", len(detail_2) == 2, len(detail_2))

# ---------- 7. generate_report 文档结构（Phase 2.5 四段式：汇总→明细→对比→分类） ----------
print("\n--- 7. generate_report 文档结构（汇总→明细→对比→分类）---")
with tempfile.TemporaryDirectory() as tmp:
    result = finance_module.generate_report(data=CAT_SAMPLE, output_dir=tmp)
    check("报告生成成功", result.get("status") == "success", result)
    path = result.get("file_path", "")
    check("文件已生成", bool(path) and os.path.exists(path), path)

    from docx import Document
    doc = Document(path)
    joined = "\n".join(p.text for p in doc.paragraphs)
    check("含「一、收支汇总」标题", "一、收支汇总" in joined, joined)
    check("含「二、交易明细」标题", "二、交易明细" in joined, joined)
    check("含「三、期间对比」标题", "三、期间对比" in joined, joined)
    check("含「四、分类统计」标题", "四、分类统计" in joined, joined)
    # 标题顺序：汇总 → 明细 → 对比 → 分类
    _idx = [joined.index(k) for k in ("一、收支汇总", "二、交易明细", "三、期间对比", "四、分类统计")]
    check("标题顺序正确（汇总→明细→对比→分类）", _idx == sorted(_idx), _idx)
    check("汇总含 总支出/总收入/结余", all(k in joined for k in ("总支出", "总收入", "结余")), joined)
    check("整篇文档无 emoji（Word 方块兼容）", not _has_emoji(joined), joined)

    check("报告含 3 张数据表（明细/对比/分类）", len(doc.tables) == 3, len(doc.tables))
    detail_table = doc.tables[0]
    check("明细表表头为 日期/时段/分类/来源/金额",
          [c.text for c in detail_table.rows[0].cells] == ["日期", "时段", "分类", "来源", "金额"],
          [c.text for c in detail_table.rows[0].cells])
    check("明细表行数 = 表头 + 4 条记录（含收入）", len(detail_table.rows) == 5, len(detail_table.rows))
    comp_table = doc.tables[1]
    check("对比表：表头 + 周/月 2 行", len(comp_table.rows) == 3, len(comp_table.rows))
    check("对比表含「本周 vs 上周」「本月 vs 上月」",
          [r.cells[0].text for r in comp_table.rows[1:]] == ["本周 vs 上周", "本月 vs 上月"],
          [r.cells[0].text for r in comp_table.rows[1:]])
    cat_table = doc.tables[2]
    check("分类表：表头 + 3 分类行", len(cat_table.rows) == 4, len(cat_table.rows))
    check("分类表表头为 分类/金额/占比",
          [c.text for c in cat_table.rows[0].cells] == ["分类", "金额", "占比"],
          [c.text for c in cat_table.rows[0].cells])

# ---------- 8. Phase 2.5 默认范围与 scope 参数 ----------
print("\n--- 8. 默认范围（今日/近3天）与 scope 参数 ---")
_TODAY = datetime.now().strftime("%Y-%m-%d")
_D1 = (datetime.now() - timedelta(days=1)).strftime("%Y-%m-%d")
_D2 = (datetime.now() - timedelta(days=2)).strftime("%Y-%m-%d")
_D5 = (datetime.now() - timedelta(days=5)).strftime("%Y-%m-%d")
SCOPE_SAMPLE = [
    {"id": "t1", "type": "expense", "date": _TODAY, "amount": 100, "source": "饭", "category": "餐饮"},
    {"id": "t2", "type": "income", "date": _TODAY, "amount": 500, "source": "工资", "category": "收入"},
    {"id": "d1", "type": "expense", "date": _D1, "amount": 20, "source": "打车", "category": "交通"},
    {"id": "d2", "type": "expense", "date": _D2, "amount": 30, "source": "奶茶", "category": "饮品"},
    {"id": "old", "type": "expense", "date": _D5, "amount": 999, "source": "购物", "category": "购物"},
]

# _get_default_date_range：今天有记录 → 今天
rd, s3, e3 = finance_module._get_default_date_range(SCOPE_SAMPLE)
check("默认范围报告日为今天", rd == _TODAY, (rd, _TODAY))
check("近三天起点 = 今天-2", s3 == (datetime.now() - timedelta(days=2)).strftime("%Y-%m-%d"), s3)
check("近三天终点 = 今天", e3 == _TODAY, e3)
# 今天无记录 → 最近有数据日
rd2, s32, e32 = finance_module._get_default_date_range(
    [{"id": "a", "type": "expense", "date": _D1, "amount": 5}])
check("今天无记录时取最近有数据日", rd2 == _D1, rd2)

with tempfile.TemporaryDirectory() as tmp:
    # default：汇总=今日(支出100/收入500)，明细=近3天，不含5天前
    res = finance_module.generate_report(data=SCOPE_SAMPLE, output_dir=tmp)
    check("default 返回 scope=default", res.get("scope") == "default", res)
    check("default 汇总仅含今日数据（支出100/收入500/结余400）",
          res["summary"] == {"income_total": 500, "expense_total": 100, "balance": 400},
          res["summary"])
    doc = Document(res["file_path"])
    det_dates = [r.cells[0].text for r in doc.tables[0].rows[1:]]
    check("default 明细仅含近3天（不含5天前记录）",
          set(det_dates) <= {_TODAY, _D1, _D2} and _D5 not in det_dates, det_dates)
    check("default 明细条数 ≤ 30", len(det_dates) <= 30, len(det_dates))
    joined = "\n".join(p.text for p in doc.paragraphs)
    check("default 标题含「（今日）」", "（今日）" in joined, joined)
    # 分类 = 今日+近3天支出（餐饮100/交通20/饮品30），不含购物999
    cat_3d = [r.cells[0].text for r in doc.tables[2].rows[1:]]
    check("default 分类不含5天前「购物」", "购物" not in cat_3d, cat_3d)

    # all：汇总=全部(支出1149/收入500)，明细含5天前
    res_all = finance_module.generate_report(data=SCOPE_SAMPLE, output_dir=tmp, scope="all")
    check("scope=all 返回 scope=all", res_all.get("scope") == "all", res_all)
    check("scope=all 汇总含全部数据（支出1149/收入500）",
          res_all["summary"] == {"income_total": 500, "expense_total": 1149, "balance": -649},
          res_all["summary"])
    det_all = [r.cells[0].text for r in Document(res_all["file_path"]).tables[0].rows[1:]]
    check("scope=all 明细含5天前记录", _D5 in det_all, det_all)

    # today：严格今日，不含昨日/前日
    res_today = finance_module.generate_report(data=SCOPE_SAMPLE, output_dir=tmp, scope="today")
    check("scope=today 汇总仅今日（支出100/收入500/结余400）",
          res_today["summary"] == {"income_total": 500, "expense_total": 100, "balance": 400},
          res_today["summary"])
    det_today = [r.cells[0].text for r in Document(res_today["file_path"]).tables[0].rows[1:]]
    check("scope=today 明细仅今日", set(det_today) == {_TODAY}, det_today)

    # week / month：窗口正确
    res_week = finance_module.generate_report(data=SCOPE_SAMPLE, output_dir=tmp, scope="week")
    check("scope=week 返回 scope=week", res_week.get("scope") == "week", res_week)
    res_month = finance_module.generate_report(data=SCOPE_SAMPLE, output_dir=tmp, scope="month")
    check("scope=month 返回 scope=month", res_month.get("scope") == "month", res_month)
    # 非法 scope 回退 default
    res_bad = finance_module.generate_report(data=SCOPE_SAMPLE, output_dir=tmp, scope="bogus")
    check("非法 scope 回退 default", res_bad.get("scope") == "default", res_bad)

print("\n" + "=" * 60)
print(f"结果: {passed} 通过, {failed} 失败")
print("=" * 60)
sys.exit(1 if failed else 0)
